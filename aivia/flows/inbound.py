"""The two inbound doors — extracts to kg1_intake, estate to the mapper.

Thin by design: validation and lifecycle live in the layers' single
writers; these flows sequence the doors and return reports. Estate
conservation (E5): acquired u counted-excluded = every file in scope —
an unsupported dialect lands as a counted exclusion, never parsed,
never silent.
"""
import json
import pathlib
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Set

from aivia.graph import kg1_intake, kg2_mapper, kg2_translator

# THE AI-GENERATED PREFIX (Sunny's ruling 2026-09-18, Brief_M7:
# "always add this phrase 'AI-generated: '", scope "all
# agent-authored"): agent-authored description text lands on
# nodes wearing this label; approved artifact texts stay pure.
AI_PREFIX = "AI-generated: "


def receive_extract(store, reg: Dict[str, Any],
                    snap: "kg1_intake.ExtractSnapshot",
                    known_packs: Set[str]) -> "kg1_intake.IntakeReport":
    kg1_intake.validate_extract(reg, snap, known_packs)
    return kg1_intake.apply_extract(store, reg, snap)


def render_intake_report(reg: Dict[str, Any], extract_reports,
                         estate_report, read) -> str:
    """The DBA-facing rendering of the intake result data (SOP: what
    loaded, what was counted as missing, anything quarantined; a
    refusal would have named its rule before this ever rendered).
    DATA FIRST — every number here is recomputed from the graph and
    the reports, never hand-written."""
    lines = ["AIVIA INTAKE REPORT",
             f"Registered db: {reg['db_name']} (server {reg['server']}); "
             f"sources: {', '.join(reg['registered_sources'])}; "
             f"DBA: {reg['dba_team']}", ""]
    tables = read.nodes("table")
    columns = read.nodes("column")
    for rep in extract_reports:
        src = rep.source
        n_tables = sum(1 for t in tables if t.identity.startswith(f"{src}|"))
        n_cols = sum(1 for c in columns if c.identity.startswith(f"{src}|"))
        lines.append(f"[extract: {src}] loaded {n_tables} tables, "
                     f"{n_cols} columns; checks: "
                     f"{rep.check_outcomes.get('INTAKE-0..10', '?')}")
        grain_gap = rep.gap_lists.get("grain_not_declared", [])
        lines.append(f"  counted missing: {len(grain_gap)} table(s) with "
                     "no declared grain"
                     + (f" ({', '.join(grain_gap)})" if grain_gap else ""))
        for alert in rep.dba_alerts:
            lines.append(f"  QUARANTINED (needs your confirmation): {alert}")
        for decl in rep.illegal_declarations:
            lines.append(f"  refused declaration: {decl}")
        for pend in rep.pending_references:
            lines.append(f"  pending reference: {pend}")
    lines.append("")
    lines.append(f"[estate] acquired {len(estate_report.acquired)} file(s): "
                 f"{', '.join(sorted(estate_report.acquired))}")
    for exc in estate_report.counted_excluded:
        lines.append(f"  excluded (counted): {exc['file']} — {exc['reason']}")
    unresolved = [(t['name'], r)
                  for t in estate_report.trees.values()
                  for r in t['resolution_census']['unresolved']]
    for fname, ref in unresolved:
        lines.append(f"  unresolved reference: {ref} in {fname} — not in "
                     "the dictionary (counted, never guessed)")
    return "\n".join(lines)


def write_intake_result_tables(out_dir, reg, extract_reports,
                               estate_report, read) -> List[str]:
    """Contract §8: the intake report is DATA FIRST — result tables
    beside the graph, from which renderings derive. One row per
    distinct finding (a DBA troubleshoots a list, not prose): CSVs
    always; an .xlsx workbook of the same sheets when openpyxl is
    present. Returns the files written."""
    import csv as _csv
    out_dir = pathlib.Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    dict_tables = {}
    for t in read.nodes("table"):
        source, schema, table = t.identity.split("|")
        dict_tables.setdefault(table.upper(), []).append(schema)

    sheets: Dict[str, List[List[str]]] = {}
    sheets["summary"] = [["item", "value"]]
    for rep in extract_reports:
        sheets["summary"] += [
            [f"{rep.source}: checks", rep.check_outcomes.get(
                "INTAKE-0..10", "?")],
            [f"{rep.source}: quarantined join groups",
             str(len(rep.quarantined_join_groups))],
            [f"{rep.source}: refused declarations",
             str(len(rep.illegal_declarations))],
            [f"{rep.source}: pending references",
             str(len(rep.pending_references))]]
    sheets["summary"] += [
        ["estate: files acquired", str(len(estate_report.acquired))],
        ["estate: files excluded (counted)",
         str(len(estate_report.counted_excluded))]]

    # The COVERAGE CHECK — every customer needs this: their SQL scanned
    # against their delivered dictionary; what falls through names the
    # next metadata to load, or a drift they are not aware of.
    unresolved: Dict[str, Dict[str, Any]] = {}
    for fname, tree in estate_report.trees.items():
        for detail in tree["resolution_census"].get(
                "unresolved_detail", []):
            row = unresolved.setdefault(
                detail["ref"], {"n": 0, "files": set(), **detail})
            row["n"] += 1
            row["files"].add(fname)
    sources_to_load: Dict[str, Dict[str, Any]] = {}
    sheets["unresolved_references"] = [
        # literal: frame
        ["reference", "kind", "occurrences", "diagnosis", "files"]]
    for ref in sorted(unresolved, key=lambda r: -unresolved[r]["n"]):
        row = unresolved[ref]
        bare = ref.split(".")[-1].upper()
        if row["kind"] == "table" and bare in dict_tables:
            diagnosis = ("SCHEMA MISMATCH — table exists in the "
                         "dictionary under schema "
                         f"'{'/'.join(dict_tables[bare])}'; the SQL "
                         "qualifies it differently")
        elif row["kind"] == "table":
            schema = row.get("schema", "(unqualified)")
            diagnosis = (f"TABLE NOT IN DELIVERED METADATA — schema "
                         f"'{schema}' has no registered extract "
                         "covering it; see sources_to_load")
            agg = sources_to_load.setdefault(
                schema, {"tables": set(), "refs": 0})
            agg["tables"].add(bare)
            agg["refs"] += row["n"]
        else:
            table = row.get("table", "?")
            diagnosis = (f"COLUMN NOT IN DELIVERED METADATA — table "
                         f"'{table}' is delivered but carries no such "
                         "column: dictionary metadata incomplete, OR "
                         "the SQL drifted from the source (a report "
                         "that may be failing silently)")
        sheets["unresolved_references"].append(
            [ref, row["kind"], str(row["n"]), diagnosis,
             "; ".join(sorted(row["files"]))])
    sheets["sources_to_load"] = [
        # literal: frame
        ["schema / source", "missing tables", "references",
         "recommendation"]]
    for schema in sorted(sources_to_load,
                         key=lambda s: -sources_to_load[s]["refs"]):
        agg = sources_to_load[schema]
        sheets["sources_to_load"].append(
            [schema, ", ".join(sorted(agg["tables"])), str(agg["refs"]),
             "register this schema's source and load its dictionary/"
             "catalog metadata (contract §3c: org schemas come from the "
             "database catalog when no vendor dictionary covers them)"])

    sheets["grain_gaps"] = [["table", "note"]]
    for rep in extract_reports:
        for table in rep.gap_lists.get("grain_not_declared", []):
            sheets["grain_gaps"].append(
                [table, "description carries no grain declaration"])
    sheets["quarantines_and_alerts"] = [["source", "detail"]]
    for rep in extract_reports:
        for alert in rep.dba_alerts:
            sheets["quarantines_and_alerts"].append([rep.source, alert])
        for decl in rep.illegal_declarations:
            sheets["quarantines_and_alerts"].append([rep.source, decl])
    sheets["excluded_files"] = [["file", "reason"]]
    for exc in estate_report.counted_excluded:
        sheets["excluded_files"].append([exc["file"], exc["reason"]])

    written = []
    for name, rows in sheets.items():
        path = out_dir / f"{name}.csv"
        with open(path, "w", newline="") as f:
            _csv.writer(f).writerows(rows)
        written.append(str(path))
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name, rows in sheets.items():
            ws = wb.create_sheet(name[:31])
            for row in rows:
                ws.append(row)
            for cell in ws[1]:
                cell.font = openpyxl.styles.Font(bold=True)
        xlsx = out_dir / "intake_report.xlsx"
        wb.save(xlsx)
        written.append(str(xlsx))
    except ImportError:
        pass  # CSVs are the data of record; the workbook is a rendering
    (out_dir / "intake_report.txt").write_text(
        render_intake_report(reg, extract_reports, estate_report, read))
    written.append(str(out_dir / "intake_report.txt"))
    return written


@dataclass
class EstateReport:
    acquired: List[str] = field(default_factory=list)
    counted_excluded: List[Dict[str, str]] = field(default_factory=list)
    trees: Dict[str, Dict[str, Any]] = field(default_factory=dict)
    twins: Dict[str, Dict[str, Any]] = field(default_factory=dict)
    reused: List[str] = field(default_factory=list)


def receive_estate(store, reg: Dict[str, Any], estate_dir,
                   kg1_changed=None) -> EstateReport:
    estate_dir = pathlib.Path(estate_dir)
    manifest = json.loads((estate_dir / "manifest.json").read_text())
    location = manifest["location"]
    report = EstateReport()
    for path in sorted(estate_dir.rglob("*")):
        if path.is_dir() or path.name == "manifest.json":
            continue
        name = path.relative_to(estate_dir).as_posix()
        if path.suffix != ".sql":
            reason = (f"unsupported-dialect ({path.suffix.lstrip('.')} "
                      "placeholder)")
            kg2_mapper.record_exclusion(store, name, reason,
                                        manifest["as_of"])
            report.counted_excluded.append({"file": name,
                                            "reason": reason})
            continue
        try:
            tree = kg2_mapper.apply_file(
                store, reg, file_id=f"{location}{name}",
                file_name=name, text=path.read_text(),
                as_of=manifest["as_of"],
                default_schema=manifest.get("default_schema"),
                kg1_changed=kg1_changed)
        except ValueError as err:
            # conservation: a parse failure is a COUNTED exclusion with
            # the parser's own message — never fatal, never silent
            reason = f"parse-error: {err}"
            kg2_mapper.record_exclusion(store, name, reason,
                                        manifest["as_of"])
            report.counted_excluded.append({"file": name,
                                            "reason": reason})
            continue
        report.acquired.append(name)
        report.trees[name] = tree
        if tree.pop("_reused", False):
            # unchanged file, no KG1 ripple: tree AND twin stand —
            # the file quantum skips the whole chain (change quanta)
            report.reused.append(name)
            twin_node = next(
                (n for n in store.current_nodes("meaning_twin")
                 if n.identity == f"twin::{location}{name}"), None)
            if twin_node is not None:
                report.twins[name] = twin_node.properties["twin"]
                continue
        # Phase B (ADR 0077): parse and translate in the SAME RUN,
        # atomic — the twin regenerates with its tree at the file
        # quantum; the homomorphism law is asserted inside translate()
        report.twins[name] = kg2_translator.apply_twin(
            store, f"{location}{name}", tree, manifest["as_of"])
    _store_scope_descriptions(store, manifest["as_of"])
    report.join_layer = _store_join_layer(store, manifest["as_of"])
    # M5-1 (ruled: "build order is the reverse. after scopes"):
    # statements build AFTER scopes, BEFORE conditions — their
    # birth edges point DOWN at scopes; statement-rooted
    # predicates then find their parent standing.
    report.statement_layer = _store_statement_layer(
        store, manifest["as_of"])
    report.condition_layer = _store_condition_layer(
        store, manifest["as_of"])
    report.derived_layer = _store_derived_column_layer(
        store, manifest["as_of"])
    # M6 (approved 2026-09-17): the file layer runs LAST — its
    # downward edges and its R13 catch-all read every layer below
    report.file_layer = _store_file_layer(store, manifest["as_of"])
    return report


def _store_scope_descriptions(store, as_of) -> None:
    """THE SHAPE CONTRACT, M2 (Sunny's bottom-up ruling,
    2026-09-10): every scope node carries its description STORED —
    the lead render, composed from KG1 words. The verbatim law is a
    standing test: stored == recomputed (test_scope_layer)."""
    from aivia.flows import produce
    from aivia.graph.read_api import ReadApi
    from aivia.lenses import decisions
    read = ReadApi(store)
    by_id = {n.identity: n for n in read.nodes("scope")}
    for key, tree in sorted(read.trees().items()):
        for scope in decisions.named_scopes(tree):
            node = by_id.get(scope["name_key"])
            if node is None:
                continue
            lead = produce._scope_lead(read, tree, scope)
            if node.properties.get("description") == lead:
                continue
            store.append_node(
                "scope", node.identity,
                {**node.properties, "description": lead},
                as_of, node.extract_id)


SAME_TREE = "SAME-TREE scope "


def _join_targets(pred) -> list:
    """Ordered unique side targets of one ON predicate: table ids
    and same-tree scope keys, in walk order (subject before
    comparand — the ruled syntactic side order)."""
    out = []

    def walk(d):
        if isinstance(d, dict):
            r = d.get("resolves_to")
            if isinstance(r, str):
                if r.count("|") == 3:
                    t = r.rsplit("|", 1)[0]
                    if t not in out:
                        out.append(t)
                elif r.startswith(SAME_TREE):
                    s = r[len(SAME_TREE):]
                    if s not in out:
                        out.append(s)
            for v in d.values():
                walk(v)
        elif isinstance(d, list):
            for v in d:
                walk(v)
    walk(pred)
    return out


def _join_entries(scope) -> list:
    """Every ON predicate under this named scope's own subtree
    (anonymous descendants attach here per A4; nested NAMED scopes
    are their own named_scopes entries and never nest in dicts)."""
    out = []

    def walk(d):
        if isinstance(d, dict):
            for e in (d.get("join_on") or []):
                out.append(e)
            for v in d.values():
                walk(v)
        elif isinstance(d, list):
            for v in d:
                walk(v)
    walk(scope)
    return out


def _disp(target: str) -> str:
    return target.rsplit("|", 1)[-1] if "|" in target \
        else target.rsplit("::", 1)[-1]


def join_render(targets, fragment) -> str:
    """The join node's stored description — deterministic; the
    verbatim law holds stored == this recompute."""
    if len(targets) >= 2:
        return (f"Joins {_disp(targets[0])} with "
                f"{_disp(targets[1])} on {fragment}.")
    if len(targets) == 1:
        return f"Joins {_disp(targets[0])} on {fragment}."
    return f"Join on {fragment}."


def read_render(target) -> str:
    """The direct_read node's stored description — deterministic;
    the verbatim law holds stored == this recompute."""
    return f"Reads {_disp(target)}."


def _store_join_layer(store, as_of) -> dict:
    """THE FROM-STRUCTURE FAMILY, era 3 (Design_Graph_Engine,
    ratified 2026-09-14, superseding the M2 remainder rule): every
    FROM clause walks the same sided shape — scope—has_part→
    (join | direct_read)—left_side/right_side→table-or-scope. A
    join carries the OBSERVED pair + ON meaning (never a
    table→table edge; joins_to stays dictionary-only); a
    direct_read is the single-table FROM — ONE left_side, no ON,
    no type: absence lives in the KIND, never a null endpoint (it
    twins ScriptDom's NamedTableReference). The era-2 `reads`
    edge is RETIRED. Conservation returned, never silent: joins ⊎
    one_sided ⊎ no_sided == every ON entry; direct_reads counted;
    the one invariant: per scope, side-targets == read-set."""
    from aivia.graph.read_api import ReadApi
    from aivia.lenses import decisions
    read = ReadApi(store)
    by_id = {n.identity: n for n in read.nodes("scope")}
    # literal: shape
    counts = {"joins": 0, "two_sided": 0, "one_sided": 0,
              "no_sided": 0, "overflow_3plus": 0,
              "side_edges_table": 0, "side_edges_scope": 0,
              "direct_reads": 0, "reads_covered": 0}
    for key, tree in sorted(read.trees().items()):
        for scope in decisions.named_scopes(tree):
            node = by_id.get(scope["name_key"])
            if node is None:
                continue
            if any(e.from_id == node.identity
                   for e in store.current_edges("has_part")
                   if "::join#" in e.to_id or "::read#" in e.to_id):
                continue  # already materialized this boot
            side_tables = set()
            for i, pred in enumerate(_join_entries(scope), 1):
                jid = f"{node.identity}::join#{i}"
                frag = (pred.get("evidence") or {}).get("fragment", "")
                targets = _join_targets(pred)
                if len(targets) > 2:
                    counts["overflow_3plus"] += 1
                sides = targets[:2]
                counts["joins"] += 1
                counts[("two_sided" if len(sides) == 2 else
                        "one_sided" if len(sides) == 1
                        else "no_sided")] += 1
                store.append_node(
                    "join", jid,
                    # literal: shape
                    {"name": f"join#{i}",
                     "description": join_render(sides, frag),
                     "on": frag,
                     "joinType": str(pred.get("join_type") or "")},
                    as_of, node.extract_id)
                store.append_edge("has_part", node.identity, jid,
                                  {}, as_of, node.extract_id)
                for side_label, target in zip(
                        ("left_side", "right_side"), sides):
                    store.append_edge(side_label, jid, target,
                                      {}, as_of, node.extract_id)
                    if "|" in target:
                        counts["side_edges_table"] += 1
                        side_tables.add(target)
                    else:
                        counts["side_edges_scope"] += 1
            reads = []

            def walk_reads(d):
                if isinstance(d, dict):
                    for fr in (d.get("from_refs") or []):
                        r = fr.get("resolves_to")
                        if isinstance(r, str) and r.count("|") == 2:
                            reads.append(r)
                    for v in d.values():
                        walk_reads(v)
                elif isinstance(d, list):
                    for v in d:
                        walk_reads(v)
            walk_reads(scope)
            nth = 0
            for t in sorted(set(reads)):
                if t in side_tables:
                    counts["reads_covered"] += 1
                    continue  # travels through a join side
                counts["direct_reads"] += 1
                nth += 1
                rid = f"{node.identity}::read#{nth}"
                store.append_node(
                    "direct_read", rid,
                    # literal: shape
                    {"name": f"read#{nth}",
                     "description": read_render(t)},
                    as_of, node.extract_id)
                store.append_edge("has_part", node.identity, rid,
                                  {}, as_of, node.extract_id)
                store.append_edge("left_side", rid, t,
                                  {}, as_of, node.extract_id)
    return counts


# literal: schema-mirror kg2_kind_library roles
ROLE_KEYS = ("subject", "comparand", "lower_bound", "upper_bound",
             "pattern", "escape", "quantifier")


def condition_render(pred, voice) -> str:
    """The condition node's stored description — the ratified
    grammar (_voice_predicate) for leaves; a short structural
    sentence for AND/OR/NOT containers (their children carry the
    detail); evidence-fragment fallback where the grammar has no
    phrase. The verbatim law holds stored == this recompute."""
    from aivia.flows import produce
    kind = pred.get("kind", "")
    kids = pred.get("children") or []
    # literal: grammar
    if kind == "NOT" and kids:
        # GRAMMAR 2.5.0 (Sunny's ruling 2026-09-12): NOT FOLDS
        # INTO ITS CHILD — the store node speaks ONE meaning
        try:
            phrase = produce._voice_predicate(pred, voice)
        except (KeyError, TypeError, AttributeError):
            phrase = ""
        if phrase:
            return phrase
        return "The inner condition does not hold."
    # AND/OR keep the structural sentence — STRUCTURE NEVER ROWS
    # (same ruling): they frame delivery, their children carry
    # the detail
    # literal: grammar
    if kind in ("AND", "OR", "NOT"):
        # literal: grammar
        word = {"AND": f"All {len(kids)} of its parts hold.",
                "OR": f"Any of its {len(kids)} parts holds.",
                "NOT": "The inner condition does not hold."}
        return word[kind]
    try:
        phrase = produce._voice_predicate(pred, voice)
    except (KeyError, TypeError, AttributeError):
        phrase = ""  # a kind the grammar has no phrase for
    if phrase:
        # 2.7.0: the SQL author's trailing note reaches the store
        # phrase (R8 attribution — 'MED_ROUTE_CODE = 11
        # -- intravenous' was voicing a bare magic number)
        note = pred.get("annotation")
        if note and note.lower() not in phrase.lower():
            phrase = phrase.rstrip(".") + f" (noted '{note}')."
        return phrase
    frag = (pred.get("evidence") or {}).get("fragment", "")
    return f"Condition: {frag}."


def _condition_refs(pred):
    """(role, target) pairs for THIS predicate's own expressions —
    role-keyed subtrees walked through function args, stopping at
    selection interiors (their conditions are their own roots) and
    at child predicates (children carry their own refs)."""
    out = []

    def walk(d, role):
        if isinstance(d, dict):
            if "scope" in d:
                return  # selection interior — its own world
            k = d.get("kind")
            if k == "column_ref" and isinstance(
                    d.get("resolves_to"), str) \
                    and d["resolves_to"].count("|") == 3:
                out.append((role, d["resolves_to"]))
            elif k == "parameter_ref" and d.get("ref"):
                out.append((role, "@" + d["ref"].lstrip("@")))
            for v in d.values():
                walk(v, role)
        elif isinstance(d, list):
            for v in d:
                walk(v, role)
    for role in ROLE_KEYS:
        if role in pred:
            walk(pred[role], role)
    for e in pred.get("comparand_list") or []:
        walk(e, "comparand")
    return out


def _composed_condition(pred, voice) -> str:
    """R11's composed form, shared with R13 (checkpoint rulings
    C1+C2, Sunny "all three" 2026-09-17): FULLY RECURSIVE — a
    composite at ANY depth speaks its children (a nested OR reads
    "either X or Y"), never the structural summary; DEGENERATE
    leaves (the author's 1=1 idiom, flagged by the condition
    layer) are PRUNED from the composition — they stay in the
    graph, counted, but say nothing at file/statement grain."""
    from aivia.lenses import decisions

    def phrase(p_, top):
        kind = p_.get("kind", "")
        kids = p_.get("children") or []
        if p_.get("node") == "predicate" or not kids                 or kind not in ("OR", "AND"):
            if p_.get("node") == "predicate"                     and decisions.is_degenerate(p_):
                return ""
            return condition_render(p_, voice).rstrip(".")
        sub = [phrase(k, False) for k in kids]
        sub = [s for s in sub if s]
        if not sub:
            return ""
        if len(sub) == 1:
            return sub[0]
        lowered = [s[0].lower() + s[1:] for s in sub]
        if kind == "AND":
            return " and ".join(lowered)
        joined = " or ".join(lowered)
        return joined if top else "either " + joined

    text = phrase(pred, True)
    if not text:
        return ""
    return text + "."


def _render_technical_definition(read, file_id: str) -> str:
    """R13 THE CATCH-ALL (DRAFT, Brief_M6_File_Layer — Sunny
    gap-checks the real render before ratification). The file's
    technical definition, composed FROM THE GRAPH'S OWN ROWS:
    Presents (the delivery's output list — passthroughs by folded
    name, computed outputs by their stored R12 phrases) +
    Population filters (the delivery chain's WHERE-rooted
    condition phrases; CASE whens excluded — projection logic,
    not population) + Inner joins (joinType Inner on the chain;
    outer joins excluded — they do not restrict the population).
    Deterministic; verbatim law: stored == recomputed."""
    from aivia.flows import produce
    from aivia.lenses import decisions
    tree = read.trees().get(file_id)
    if tree is None:
        return ""
    store = read._store
    scopes = {sc["name_key"]: sc
              for sc in decisions.named_scopes(tree)}
    deliveries = [st["scope"] for st in tree["statements"]
                  if st.get("emits") and st.get("scope")
                  and st["scope"].get("name_key")]
    voice = produce._Voice(read, tree)

    # section 1 — Presents: the delivery's TOP-LEVEL output list
    # (nested subquery members are interior machinery, not
    # presented columns — the checkpoint's first find). The dcol
    # seq walks ALL members (the store's id grain); the top-level
    # filter applies to what is VOICED.
    dnodes = {n.identity: n
              for n in store.current_nodes("derived_column")}
    presents: List[str] = []
    for d in deliveries:
        top = [m for m in d.get("projection") or []
               if isinstance(m, dict)
               and m.get("node") == "projection_member"]
        top_ids = {id(m) for m in top}
        seq = 0
        for m in _derived_members(d):
            if _is_derived_node(m):
                seq += 1
                if id(m) not in top_ids:
                    continue
                node = dnodes.get(f"{d['name_key']}::dcol#{seq}")
                if node is not None:
                    presents.append(str(
                        node.properties.get("description")
                        or "").rstrip("."))
                    continue
            if id(m) not in top_ids:
                continue
            if m.get("star"):
                src = " and ".join(
                    produce._spoken_selection(
                        str(ref.get("table_ref") or ""))
                    for ref in d.get("from_refs") or []
                    if ref.get("table_ref")) or "its source"
                presents.append(f"every column of the {src} "
                                "selection")
                continue
            name = m.get("name")
            if name:
                presents.append(produce._readable_name(name))

    # the delivery chain — the R10 spine walk
    on_chain: List[str] = []

    def walk(scope):
        for ref in decisions.nested_sources(scope):
            rt = str(ref.get("resolves_to") or "")
            if rt.startswith("SAME-TREE scope "):
                nk = rt.replace("SAME-TREE scope ", "")
                if nk in scopes and nk not in on_chain:
                    on_chain.append(nk)
                    walk(scopes[nk])
    for d in deliveries:
        walk(d)
    chain = [d["name_key"] for d in deliveries] + on_chain

    # section 2a — Population filters, GROUPED by selection
    # (checkpoint ruling C3, "all three"): each chain scope's
    # filters open with its spoken name — 58 items gain addresses
    filters: List[str] = []
    for sk in chain:
        wheres: List[dict] = []

        def collect(d):
            if isinstance(d, dict):
                if isinstance(d.get("where"), dict):
                    wheres.append(d["where"])
                for v in d.values():
                    collect(v)
            elif isinstance(d, list):
                for v in d:
                    collect(v)
        collect(scopes.get(sk) or next(
            (x for x in deliveries if x["name_key"] == sk), {}))
        items = []
        for pred in wheres:
            txt = _composed_condition(pred, voice).rstrip(".")
            if txt:
                items.append(txt[0].lower() + txt[1:])
        if items:
            spoken = produce._spoken_selection(
                sk.rsplit("::", 1)[-1])
            filters.append(f"In the {spoken} selection: "
                           + "; ".join(items) + ".")

    # section 2b — Inner joins on the chain (stored phrases)
    joins = {n.identity: n for n in store.current_nodes("join")}
    inner: List[str] = []
    for sk in chain:
        jids = sorted(
            (j for j in joins if j.startswith(f"{sk}::join#")),
            key=lambda j: int(j.rsplit("#", 1)[-1]))
        for jid in jids:
            n = joins[jid]
            if str(n.properties.get("joinType") or "") == "Inner":
                inner.append(str(n.properties.get("description")
                                 or "").rstrip("."))

    parts: List[str] = []
    if presents:
        parts.append("Presents: " + "; ".join(presents) + ".")
    if filters:
        parts.append("Population filters: " + " ".join(filters))
    if inner:
        parts.append("Inner joins: " + "; ".join(inner) + ".")
    return " ".join(parts)


def _store_file_layer(store, as_of) -> dict:
    """M6 THE FILE LAYER (Brief_M6_File_Layer, Sunny's 'yes, yes,
    yes. approved' 2026-09-17): file—has_part→statement for EVERY
    statement — THE 31-STATEMENT DEBT RETIRES at its named
    landing step — + file—has_part→param; the TWO GOVERNANCE
    FIELDS land on the node: technical_definition (R13, verbatim)
    and description (the approved Scribe summary ONLY)."""
    from aivia.graph.read_api import ReadApi
    read = ReadApi(store)
    # literal: shape
    counts = {"files": 0, "statement_edges": 0, "param_edges": 0,
              "definitions": 0}
    files = {n.identity: n for n in store.current_nodes("file")}
    st_ids = {n.identity for n in read.nodes("statement")}
    p_ids = {n.identity for n in read.nodes("param")}
    for key, tree in sorted(read.trees().items()):
        fnode = files.get(key)
        if fnode is None:
            continue
        counts["files"] += 1
        already = any(e.from_id == key and e.to_id in st_ids
                      for e in store.current_edges("has_part"))
        if already:  # this boot already materialized the layer
            continue
        for i in range(len(tree["statements"])):
            sid = f"{key}::stmt/{i + 1}"
            if sid in st_ids:
                store.append_edge("has_part", key, sid, {},
                                  as_of, fnode.extract_id)
                counts["statement_edges"] += 1
        for pid in sorted(p_ids):
            if pid.startswith(f"{key}::param/"):
                store.append_edge("has_part", key, pid, {},
                                  as_of, fnode.extract_id)
                counts["param_edges"] += 1
        td = _render_technical_definition(read, key)
        if td:
            props = dict(fnode.properties)
            props["technical_definition"] = td
            counts["definitions"] += 1
            store.append_node("file", key, props, as_of,
                              fnode.extract_id)
    return counts


def _twin_statement_subkinds(store, tree) -> Dict[int, str]:
    """M5-3 (resolved by investigation, Sunny shown 2026-09-17):
    subkind is READ from the translator's T-2 field on the twin —
    never re-derived here (one writer since 2026-09-06)."""
    for mt in store.current_nodes("meaning_twin"):
        twin = mt.properties.get("twin", {})
        if twin.get("file") != tree.get("name"):
            continue
        out = {}
        for n in twin["nodes"]:
            m = re.fullmatch(r"/statements/(\d+)",
                             n.get("points_at", ""))
            if m and n.get("kind") == "statement" \
                    and n.get("subkind"):
                out[int(m.group(1))] = n["subkind"]
        return out
    return {}


def _render_statement_descriptions(read) -> Dict[str, str]:
    """The R11 recompute surface — the verbatim law's second
    reading (stored == recomputed, byte-exact). Voiced statements
    only; the silent kinds return nothing here by construction."""
    from aivia.flows import produce

    out: Dict[str, str] = {}
    for key, tree in sorted(read.trees().items()):
        voice = produce._Voice(read, tree)
        for i, st in enumerate(tree["statements"]):
            phrase = None
            if st.get("predicate") is not None:
                phrase = _composed_condition(st["predicate"], voice)
            text = produce.statement_phrase(
                st, predicate_phrase=phrase)
            if text:
                out[f"{key}::stmt/{i + 1}"] = text
    return out


def _store_statement_layer(store, as_of) -> dict:
    """M5 THE STATEMENT LAYER (Brief_M5_Statement_Layer, Sunny's
    'approved' 2026-09-17): one node per statement — identity
    file::stmt/<position> (the 2026-09-10 ruling) — built after
    scopes, before conditions (M5-1). Descriptions: R11 renders
    for the voiced kinds; operational statements store NOTHING
    (the (b) ruling — the emptiness COUNTED); an unlisted
    data-producing kind is a COUNTED remainder, never silent.
    statement—has_part→scope birth edges point DOWN at verified
    scopes; the 31 operational stay counted-missing until M6."""
    from aivia.graph.read_api import ReadApi
    read = ReadApi(store)
    files = {n.identity: n for n in store.current_nodes("file")}
    scope_ids = {n.identity for n in read.nodes("scope")}
    texts = _render_statement_descriptions(read)
    # literal: shape
    counts = {"statements": 0, "voiced": 0, "operational": 0,
              "scope_edges": 0, "statement_remainders": {}}
    for key, tree in sorted(read.trees().items()):
        existing = [n for n in read.nodes("statement")
                    if n.identity.startswith(f"{key}::stmt/")]
        if existing:  # already materialized this boot
            counts["statements"] += len(existing)
            continue
        extract = files[key].extract_id
        subk = _twin_statement_subkinds(store, tree)
        for i, st in enumerate(tree["statements"]):
            sid = f"{key}::stmt/{i + 1}"
            kind = st.get("statement_kind", "")
            # literal: shape
            props: Dict[str, Any] = {"name": f"stmt/{i + 1}",
                                     "does": kind}
            if i in subk:
                props["subkind"] = subk[i]
            text = texts.get(sid)
            if text:
                props["description"] = text
                counts["voiced"] += 1
            elif subk.get(i) == "operational":
                counts["operational"] += 1
            elif st.get("scope") or st.get("ctes"):
                counts["statement_remainders"][kind] = \
                    counts["statement_remainders"].get(kind, 0) + 1
            store.append_node("statement", sid, props, as_of,
                              extract)
            counts["statements"] += 1
            for cte in st.get("ctes") or []:
                nk = cte.get("name_key")
                if nk in scope_ids:
                    store.append_edge("has_part", sid, nk, {},
                                      as_of, extract)
                    counts["scope_edges"] += 1
            nk = (st.get("scope") or {}).get("name_key")
            if nk in scope_ids:
                store.append_edge("has_part", sid, nk, {},
                                  as_of, extract)
                counts["scope_edges"] += 1
    return counts


def _store_condition_layer(store, as_of) -> dict:
    """THE CONDITION LAYER, M3: every predicate under a named scope
    becomes a condition node with its stored voiced phrase — ON
    roots parent to their JOIN, where/case roots to the scope,
    nested predicates to their parent condition (kind is a
    PROPERTY, never a label). resolves_to carries the ROLE;
    parameters mint as param nodes with uses_param birth edges.
    Conservation returned, never silent."""
    from aivia.flows import produce
    from aivia.graph.read_api import ReadApi
    from aivia.lenses import decisions
    read = ReadApi(store)
    by_id = {n.identity: n for n in read.nodes("scope")}
    # literal: shape
    counts = {"conditions": 0, "roots_join": 0, "roots_scope": 0,
              "nested": 0, "degenerate": 0, "resolves_column": 0,
              "resolves_param": 0, "params": 0, "uses_param": 0,
              "held_statement_rooted": 0}
    for key, tree in sorted(read.trees().items()):
        voice = produce._Voice(read, tree)
        made_params = set()

        def ensure_param(ref, extract_id):
            pid = f"{key}::param/{ref}"
            if pid in made_params or any(
                    n.identity == pid
                    for n in read.nodes("param")):
                return pid
            made_params.add(pid)
            words = ref.lstrip("@")
            store.append_node(
                "param", pid,
                # literal: shape
                {"name": ref,
                 "description": f"Parameter {ref} of this "
                                f"procedure ({words})."},
                as_of, extract_id)
            counts["params"] += 1
            return pid

        for scope in decisions.named_scopes(tree):
            node = by_id.get(scope["name_key"])
            if node is None:
                continue
            prefix = node.identity + "::cond#"
            if any(n.identity.startswith(prefix)
                   for n in read.nodes("condition")):
                continue  # already materialized this boot
            seq = 0
            scope_params = set()

            def emit(pred, parent_id, is_root_kind):
                nonlocal seq
                seq += 1
                cid = f"{node.identity}::cond#{seq}"
                kind = pred.get("kind", "")
                leaf = pred.get("node") == "predicate"
                degen = bool(leaf and decisions.is_degenerate(pred))
                if degen:
                    counts["degenerate"] += 1
                store.append_node(
                    "condition", cid,
                    # literal: shape
                    {"name": f"cond#{seq}", "kind": kind,
                     "degenerate": "true" if degen else "false",
                     "description": condition_render(pred, voice),
                     "fragment": (pred.get("evidence") or {})
                     .get("fragment", "")},
                    as_of, node.extract_id)
                store.append_edge("has_part", parent_id, cid,
                                  {}, as_of, node.extract_id)
                counts["conditions"] += 1
                counts[is_root_kind] += 1
                for role, target in _condition_refs(pred):
                    if target.startswith("@"):
                        pid = ensure_param(target, node.extract_id)
                        store.append_edge(
                            "resolves_to", cid, pid,
                            {"role": role}, as_of, node.extract_id)
                        counts["resolves_param"] += 1
                        scope_params.add(pid)
                    else:
                        store.append_edge(
                            "resolves_to", cid, target,
                            {"role": role}, as_of, node.extract_id)
                        counts["resolves_column"] += 1
                for child in pred.get("children") or []:
                    emit(child, cid, "nested")

            for i, pred in enumerate(_join_entries(scope), 1):
                emit(pred, f"{node.identity}::join#{i}", "roots_join")
            wheres, whens = [], []

            def collect(d):
                if isinstance(d, dict):
                    if isinstance(d.get("where"), dict):
                        wheres.append(d["where"])
                    for w in d.get("whens") or []:
                        if isinstance(w.get("when"), dict):
                            whens.append(w["when"])
                    for v in d.values():
                        collect(v)
                elif isinstance(d, list):
                    for v in d:
                        collect(v)
            collect(scope)
            for pred in wheres + whens:
                emit(pred, node.identity, "roots_scope")
            for pid in sorted(scope_params):
                store.append_edge("uses_param", node.identity, pid,
                                  {}, as_of, node.extract_id)
                counts["uses_param"] += 1

        # M5 (the M5-1 ordering makes this reachable; FL8 closed —
        # Sunny "real value"): statement-rooted predicates parent
        # to their STATEMENT node; held_statement_rooted counts
        # every condition minted here — measured, never zero-by-
        # unreachability.
        st_nodes = {n.identity: n for n in read.nodes("statement")}
        for i, st in enumerate(tree["statements"]):
            pred = st.get("predicate")
            if pred is None:
                continue
            sid = f"{key}::stmt/{i + 1}"
            snode = st_nodes.get(sid)
            if snode is None:
                continue  # statements not built (pre-M5 store)
            prefix = sid + "::cond#"
            if any(n.identity.startswith(prefix)
                   for n in read.nodes("condition")):
                continue  # already materialized this boot
            st_seq = 0
            st_params: Set[str] = set()

            def emit_st(pred_, parent_id):
                nonlocal st_seq
                st_seq += 1
                cid = f"{sid}::cond#{st_seq}"
                kind = pred_.get("kind", "")
                leaf = pred_.get("node") == "predicate"
                degen = bool(leaf and decisions.is_degenerate(pred_))
                if degen:
                    counts["degenerate"] += 1
                store.append_node(
                    "condition", cid,
                    # literal: shape
                    {"name": f"cond#{st_seq}", "kind": kind,
                     "degenerate": "true" if degen else "false",
                     "description": condition_render(pred_, voice),
                     "fragment": (pred_.get("evidence") or {})
                     .get("fragment", "")},
                    as_of, snode.extract_id)
                store.append_edge("has_part", parent_id, cid,
                                  {}, as_of, snode.extract_id)
                counts["conditions"] += 1
                counts["held_statement_rooted"] += 1
                for role, target in _condition_refs(pred_):
                    if target.startswith("@"):
                        pid = ensure_param(target, snode.extract_id)
                        store.append_edge(
                            "resolves_to", cid, pid,
                            {"role": role}, as_of, snode.extract_id)
                        counts["resolves_param"] += 1
                        st_params.add(pid)
                    else:
                        store.append_edge(
                            "resolves_to", cid, target,
                            {"role": role}, as_of, snode.extract_id)
                        counts["resolves_column"] += 1
                for child in pred_.get("children") or []:
                    emit_st(child, cid)

            emit_st(pred, sid)
            for pid in sorted(st_params):
                store.append_edge("uses_param", sid, pid,
                                  {}, as_of, snode.extract_id)
                counts["uses_param"] += 1
    return counts


# The kinds whose projection members ARE derived_column nodes
# (M4's census law): the composite expression kinds + case
# (Expression_Kinds' value-branching row).
_DERIVED_OP_KINDS = frozenset(
    # literal: schema-mirror kg2_kind_library COMPOSITE_EXPR_KINDS
    ("function", "arithmetic", "unary", "cast", "case"))


def _derived_members(scope):
    """Projection members under one named scope, path-ordered —
    nested subquery projections included (they live inside this
    scope's dict; named scopes are never nested in each other).
    The mapper's `select_refs` key is SKIPPED: it aliases the same
    member expressions (documented at the mapper), and walking it
    double-counts — the 4 phantom members the M4 pre-build
    measurement caught."""
    out = []

    def walk(d):
        if isinstance(d, dict):
            if d.get("node") == "projection_member":
                out.append(d)
            for k, v in d.items():
                if k == "select_refs":
                    continue
                walk(v)
        elif isinstance(d, list):
            for v in d:
                walk(v)
    walk(scope)
    return out


def _is_derived_node(member) -> bool:
    """The sealed M4 census law: a member is a NODE when its
    defining expression is a computed kind, or it is a NAMED
    literal. Passthroughs (column_ref/star) and anonymous
    EXISTS-SELECT literals are counted, never minted."""
    kind = (member.get("expression") or {}).get("kind")
    if kind in _DERIVED_OP_KINDS:
        return True
    return kind == "literal" and bool(member.get("name"))


def _member_cites(member):
    """4-part dictionary columns referenced anywhere under this
    member's defining expression (subquery interiors included —
    the output draws on them)."""
    out = []

    def walk(d):
        if isinstance(d, dict):
            if d.get("kind") == "column_ref" and isinstance(
                    d.get("resolves_to"), str) \
                    and d["resolves_to"].count("|") == 3:
                out.append(d["resolves_to"])
            for v in d.values():
                walk(v)
        elif isinstance(d, list):
            for v in d:
                walk(v)
    walk(member.get("expression") or {})
    return out


def _store_derived_column_layer(store, as_of) -> dict:
    """THE DERIVED-COLUMN LAYER, M4 (Sunny's go 2026-09-16): every
    computed output under a named scope becomes a derived_column
    node with its stored R12 phrase — scope—has_part→derived_column
    birth edges (STAY FLAT: one node per output, the expression
    tree stays at L1) + scope—cites→column for every dictionary
    column the scope's outputs draw on (distinct pairs, store
    grain). Conservation returned, never silent."""
    from aivia.flows import produce
    from aivia.graph.read_api import ReadApi
    from aivia.lenses import decisions
    read = ReadApi(store)
    by_id = {n.identity: n for n in read.nodes("scope")}
    # literal: shape
    counts = {"derived_columns": 0, "operations": 0,
              "named_literals": 0, "passthrough_skipped": 0,
              "passthrough_renamed": 0,
              "anonymous_skipped": 0, "cites": 0,
              "function_remainders": {}}
    for key, tree in sorted(read.trees().items()):
        voice = produce._Voice(read, tree)
        for scope in decisions.named_scopes(tree):
            node = by_id.get(scope["name_key"])
            if node is None:
                continue
            prefix = node.identity + "::dcol#"
            already = any(n.identity.startswith(prefix)
                          for n in read.nodes("derived_column"))
            seq = 0
            cited = {e.to_id for e in read.edges("cites")
                     if e.from_id == node.identity}
            for member in _derived_members(scope):
                for col in _member_cites(member):
                    if col not in cited:
                        cited.add(col)
                        store.append_edge("cites", node.identity,
                                          col, {}, as_of,
                                          node.extract_id)
                        counts["cites"] += 1
                if not _is_derived_node(member):
                    expr = member.get("expression") or {}
                    expr_kind = expr.get("kind")
                    if expr_kind in ("column_ref", "star"):
                        counts["passthrough_skipped"] += 1
                        # the (c) trade-off, COUNTED (Sunny
                        # accepted 2026-09-16): a RENAMED
                        # passthrough (author's alias, no
                        # computation) is not a node and not an
                        # ask card — visible here, never silent
                        ref_tail = (expr.get("ref") or ""
                                    ).rsplit(".", 1)[-1]
                        if (expr_kind == "column_ref"
                                and member.get("name")
                                and ref_tail.lower()
                                != str(member["name"]).lower()):
                            counts["passthrough_renamed"] += 1
                    else:
                        counts["anonymous_skipped"] += 1
                    continue
                seq += 1
                if already:
                    continue  # this scope materialized a prior boot
                cid = f"{node.identity}::dcol#{seq}"
                expr = member.get("expression") or {}
                operation = ((expr.get("name") or "").upper()
                             if expr.get("kind") == "function"
                             else expr.get("kind", ""))
                derivation = ("named_literal"
                              if expr.get("kind") == "literal"
                              else "operation")
                counts["derived_columns"] += 1
                counts["operations" if derivation == "operation"
                       else "named_literals"] += 1
                # literal: shape
                props = {"name": member.get("name") or f"dcol#{seq}",
                         "derivation": derivation,
                         "operation": operation,
                         "description": produce.derived_phrase(
                             member, voice),
                         "fragment": (member.get("evidence") or {})
                         .get("fragment", "")}
                if member.get("position") is not None:
                    props["position"] = member["position"]
                store.append_node("derived_column", cid, props,
                                  as_of, node.extract_id)
                store.append_edge("has_part", node.identity, cid,
                                  {}, as_of, node.extract_id)
        # the remainder census rides the receipt (R12's counted law)
        for op_name, n in sorted(voice.function_remainders.items()):
            counts["function_remainders"][op_name] = \
                counts["function_remainders"].get(op_name, 0) + n
    return counts


def receive_pbi(store, pbi_dir) -> int:
    """THE CONSUMPTION LAYER (Sunny's ruling 2026-09-08: every proc
    feeds a PBI report; the mapping is real even where the shell is
    synthetic). Loads pbi_snapshot/reports.json as PBI Report nodes
    — label chosen to BE the words users say — with executes
    mappings; DISPLAYS are composed from the executed procs' twins
    (delivery projection members: what the report actually shows,
    derived, never invented)."""
    import json as _json
    import pathlib as _pl

    from aivia.graph.read_api import ReadApi
    pbi_dir = _pl.Path(pbi_dir)
    if not (pbi_dir / "reports.json").is_file():
        return 0
    manifest = _json.loads((pbi_dir / "manifest.json").read_text())
    rows = _json.loads((pbi_dir / "reports.json").read_text())
    read = ReadApi(store)
    trees = read.trees()
    rel_to_key = {}
    for key, tree in trees.items():
        rel_to_key[tree.get("name", key)] = key

    def displays_of(rel):
        tree = trees.get(rel_to_key.get(rel, rel))
        if not tree:
            return []
        out = []
        for stmt in tree.get("statements", []):
            if not stmt.get("emits") or not stmt.get("scope"):
                continue
            arms = stmt["scope"].get("combination_arms")
            shape = arms[0] if arms else stmt["scope"]
            for m in shape.get("projection", []):
                if m.get("name"):
                    out.append(m["name"])
        return out[:8]

    n = 0
    for row in rows:
        slug = row["name"].lower().replace(" ", "-")
        displays = []
        executes = []
        resolved, unresolved = [], []
        for rel in row["executes"]:
            executes.append(rel_to_key.get(rel, rel))
            displays += displays_of(rel)
            if rel in rel_to_key:
                resolved.append(rel_to_key[rel])
            else:
                unresolved.append(rel)  # COUNTED, never guessed
        rid = f"pbi://sepsis/{slug}"
        # bound_fields keyed by resolved file identity (M7-A,
        # per-proc for the Q3 concatenation)
        bound = {}
        for rel, cols in (row.get("bound_fields") or {}).items():
            if rel in rel_to_key:
                bound[rel_to_key[rel]] = cols
        store.append_node(
            "pbi_report", rid,
            # literal: shape
            {"name": row["name"],
             "description": row["description"],
             "displays": displays, "executes": executes,
             "unresolved_executes": unresolved,
             "bound_fields": bound,
             "source": row["source"]},
            manifest["as_of"], f"pbi@{manifest['as_of']}")
        # M7: executes becomes a REAL edge — resolved targets only
        for fid in resolved:
            store.append_edge("executes", rid, fid, {},
                              manifest["as_of"],
                              f"pbi@{manifest['as_of']}")
        n += 1
    # the derived pair (after every node exists; the file layer
    # ran in receive_estate, so the catch-alls stand)
    read = ReadApi(store)
    for rpt in store.current_nodes("pbi_report"):
        td = _render_report_definition(read, rpt.identity)
        if td and rpt.properties.get("technical_definition") != td:
            props = dict(rpt.properties)
            props["technical_definition"] = td
            store.append_node("pbi_report", rpt.identity, props,
                              manifest["as_of"],
                              f"pbi@{manifest['as_of']}")
    return n


# M7 (Q4 "a"): the standing disclosure — every report definition
# carries it until THE REPORT-LAYER EXTRACTION (the named future
# brief) lands the PBI-side filters
REPORT_DISCLOSURE = ("Filters shown are the procedure's; the "
                     "report may filter further in Power BI.")


def _cherry_pick_presents(file_td: str, bound_cols) -> str:
    """The Presents section filtered to the report's bound fields
    (matched through the SAME readable-name fold the render used —
    one home); the star item survives (the model's imports beyond
    named outputs are the star's columns); unbound named items
    DROP — that is the pick."""
    from aivia.flows import produce
    if "Presents: " not in file_td:
        return file_td
    head, rest = file_td.split("Presents: ", 1)
    for marker in ("Population filters: ", "Inner joins: "):
        if marker in rest:
            body, tail = rest.split(marker, 1)
            tail = marker + tail
            break
    else:
        body, tail = rest, ""
    items = [i.strip() for i in
             body.rstrip(". ").split("; ") if i.strip()]
    folded = {produce._readable_name(c) for c in bound_cols}
    kept = []
    for item in items:
        if item.startswith("every column of the "):
            kept.append(item)
            continue
        key = item.split(":", 1)[0].strip().lower()
        if key in folded:
            kept.append(item)
    if not kept:
        kept = items  # nothing matched: fall back whole, honest
    return (head + "Presents: " + "; ".join(kept) + ". " + tail)


def _render_report_definition(read, report_id: str) -> str:
    """M7 — the report's technical definition (ds.report_derived_
    pair): the executed files' catch-alls CHERRY-PICKED to the
    report's bound fields, population whole, concatenated per
    proc when plural (Q3 — labeled 'From <PROC>:'), THE
    DISCLOSURE LINE last. A shell (no bound fields) carries the
    whole catch-all. Deterministic; verbatim law."""
    store = read._store
    rpt = next((n for n in store.current_nodes("pbi_report")
                if n.identity == report_id), None)
    if rpt is None:
        return ""
    files = {n.identity: n for n in store.current_nodes("file")}
    targets = [e.to_id for e in store.current_edges("executes")
               if e.from_id == report_id and e.to_id in files]
    bound = rpt.properties.get("bound_fields") or {}
    parts = []
    for fid in targets:
        ftd = str(files[fid].properties.get(
            "technical_definition") or "")
        if not ftd:
            continue
        cols = bound.get(fid) or []
        section = (_cherry_pick_presents(ftd, cols) if cols
                   else ftd)
        if len(targets) > 1:
            base = fid.rsplit("/", 1)[-1].removesuffix(".sql")
            section = f"From {base}: {section}"
        parts.append(section.rstrip())
    if not parts:
        return ""
    return " ".join(parts) + " " + REPORT_DISCLOSURE


def receive_descriptions(store, path) -> int:
    """E1 estate-data loading: the Scribe's committed drafts
    (descriptions.json) land as kg3 description artifacts at boot —
    machine-authored, 'drafted', basis-stamped. Estate data, not a
    journal act: every boot (test or prod) speaks the same
    aboutness; blessing is a later human act in the journal."""
    import json as _json
    import pathlib as _pl

    from aivia.flows import describe as _describe
    path = _pl.Path(path)
    if not path.is_file():
        return 0
    data = _json.loads(path.read_text())
    # M7: a batch may carry MIXED statuses (per-identity
    # "statuses" overrides the file-level "status") — the carried
    # approval rides beside fresh drafts
    default = data.get("status", "drafted")
    statuses = data.get("statuses") or {}
    by_status = {}
    for ident, text in data["descriptions"].items():
        by_status.setdefault(statuses.get(ident, default),
                             {})[ident] = text
    n = 0
    for status, drafts in sorted(by_status.items()):
        n += _describe.land(store, drafts, basis=data["basis"],
                            created_at=data["created_at"],
                            status=status)
    # M6 (Sunny's "APPROVED" 2026-09-17): an APPROVED artifact's
    # text lands on the file NODE — this loader is the ONE writer
    # of file.description (it runs after the file layer and holds
    # the approval act; empty-until-approved stays the counted
    # posture). THE AI-GENERATED PREFIX (Sunny's 2026-09-18
    # ruling, scope "all agent-authored"): every text this writer
    # lands is a Scribe draft (author agent:scribe by
    # construction), so the node property carries the provenance
    # label; the approved artifact text stays pure.
    approved_idents = {i for i in data["descriptions"]
                       if statuses.get(i, default) == "approved"}
    if approved_idents:
        files = {f.identity: f for f in store.current_nodes("file")}
        files.update({r.identity: r
                      for r in store.current_nodes("pbi_report")})
        for identity, text in sorted(data["descriptions"].items()):
            if identity not in approved_idents:
                continue
            node = files.get(identity)
            if node is None or not text:
                continue
            labeled = AI_PREFIX + text
            if node.properties.get("description") == labeled:
                continue
            props = dict(node.properties)
            props["description"] = labeled
            store.append_node(node.label, identity, props,
                              data["created_at"], node.extract_id)
    return n
